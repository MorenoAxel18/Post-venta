"""Lectura de la cartera (Excel o CSV) en filas columna -> valor.

Los encabezados se normalizan: "Teléfono Celular" -> "telefono_celular".
"""
import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime

from app.comun.errores import ApiError

EXTENSIONES = (".xlsx", ".xls", ".csv")


@dataclass
class Fila:
    numero: int  # número de fila como se ve en Excel (el encabezado es la fila 1)
    valores: dict[str, str]


@dataclass
class ArchivoCartera:
    columnas: list[str]
    filas: list[Fila]


def normalizar_encabezado(original: object) -> str:
    s = str(original or "").replace("\ufeff", "").strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def _texto(valor: object) -> str:
    """Convierte el valor de una celda a texto como lo vería la persona en Excel."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # teléfonos e IDs guardados como número
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y") if valor.time() == datetime.min.time() else valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def leer(nombre_archivo: str, contenido: bytes, max_filas: int) -> ArchivoCartera:
    nombre = (nombre_archivo or "").lower()
    if not contenido:
        raise ApiError.invalido("El archivo está vacío.")
    if not nombre.endswith(EXTENSIONES):
        raise ApiError.invalido("Formato no soportado. Sube un archivo .xlsx, .xls o .csv.")

    try:
        if nombre.endswith(".csv"):
            encabezado, filas_crudas = _leer_csv(contenido)
        elif nombre.endswith(".xlsx"):
            encabezado, filas_crudas = _leer_xlsx(contenido)
        else:
            encabezado, filas_crudas = _leer_xls(contenido)
    except ApiError:
        raise
    except Exception as e:  # archivo dañado o con formato inesperado
        raise ApiError.invalido("No se pudo leer el archivo. Verifica que no esté dañado.") from e

    columnas = [normalizar_encabezado(c) for c in encabezado]
    vistos: set[str] = set()
    for c in columnas:
        if c and c in vistos:
            raise ApiError.invalido(f'La columna "{c}" está repetida en el encabezado.')
        vistos.add(c)

    filas: list[Fila] = []
    for numero, crudos in filas_crudas:
        valores = {col: _texto(crudos[i]) if i < len(crudos) else "" for i, col in enumerate(columnas) if col}
        if any(valores.values()):
            filas.append(Fila(numero, valores))
        if len(filas) > max_filas:
            raise ApiError.invalido(f"El archivo tiene más de {max_filas:,} registros. Divídelo en varias campañas.")

    if not filas:
        raise ApiError.invalido("El archivo no tiene registros debajo del encabezado.")
    return ArchivoCartera([c for c in columnas if c], filas)


def _leer_csv(contenido: bytes):
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")  # CSV guardado desde Excel en Windows
    primera = texto.split("\n", 1)[0]
    # Excel en español suele exportar CSV separado por punto y coma
    delimitador = ";" if primera.count(";") > primera.count(",") else ","
    lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
    encabezado = next(lector, [])
    return encabezado, ((i, [v.strip() for v in fila]) for i, fila in enumerate(lector, start=2))


def _leer_xlsx(contenido: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = wb.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas, None)
    if encabezado is None:
        raise ApiError.invalido("La primera hoja del archivo está vacía.")
    return list(encabezado), ((i, list(f)) for i, f in enumerate(filas, start=2))


def _leer_xls(contenido: bytes):
    import xlrd

    libro = xlrd.open_workbook(file_contents=contenido)
    hoja = libro.sheet_by_index(0)
    if hoja.nrows == 0:
        raise ApiError.invalido("La primera hoja del archivo está vacía.")

    def valor(celda):
        if celda.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate.xldate_as_datetime(celda.value, libro.datemode)
        return celda.value

    encabezado = [valor(c) for c in hoja.row(0)]
    return encabezado, ((r + 1, [valor(c) for c in hoja.row(r)]) for r in range(1, hoja.nrows))
